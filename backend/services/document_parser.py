"""Extract text and basic quality metadata from common tender documents.

PDFs containing scanned page images need OCR. This module deliberately does not
pretend that an image-only PDF was successfully read.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from hashlib import sha256
from pathlib import Path
from typing import Any


@dataclass
class DocumentExtraction:
    source_file: str
    file_type: str
    text: str
    page_count: int | None
    character_count: int
    sha256: str
    warnings: list[str]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def extract_document(file_path: str | Path) -> DocumentExtraction:
    path = Path(file_path)
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    suffix = path.suffix.lower()
    warnings: list[str] = []
    page_count: int | None = None

    if suffix in {".txt", ".md", ".text"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        page_count = None
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError("PDF support requires pypdf. Run: pip install -r requirements.txt") from exc
        reader = PdfReader(str(path))
        page_count = len(reader.pages)
        chunks = []
        pages_without_text = 0
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if not page_text:
                pages_without_text += 1
            chunks.append(f"[Page {page_number}]\n{page_text}")
        text = "\n\n".join(chunks)
        if pages_without_text:
            warnings.append(
                f"{pages_without_text} of {page_count} PDF page(s) yielded no selectable text. "
                "If these pages contain scans, OCR is required before their requirements can be extracted."
            )
        if len(text.strip()) < 80:
            warnings.append("Very little text was extracted; treat results as incomplete and inspect the PDF manually.")
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("DOCX support requires python-docx. Run: pip install -r requirements.txt") from exc
        document = Document(str(path))
        chunks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_index, table in enumerate(document.tables, start=1):
            chunks.append(f"[Table {table_index}]")
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
        page_count = None  # DOCX does not reliably expose rendered page boundaries.
    else:
        raise ValueError(f"Unsupported file type '{suffix}'. Use TXT, MD, PDF, or DOCX.")

    if not text.strip():
        warnings.append("No readable text was extracted from this document.")

    return DocumentExtraction(
        source_file=path.name,
        file_type=suffix.lstrip("."),
        text=text,
        page_count=page_count,
        character_count=len(text),
        sha256=sha256(path.read_bytes()).hexdigest(),
        warnings=warnings,
    )


def extract_text(file_path: str | Path) -> str:
    """Backward-compatible helper returning only extracted text."""
    return extract_document(file_path).text
